"""Shadow bets (what each filter blocked), second-opinion scorecard, maker fills from the price history."""
from datetime import timedelta

import pytest
from sqlalchemy import select

from backend.betting import orders, portfolio, shadow
from backend.config import settings
from backend.db.database import SessionLocal
from backend.db.models import Market, MarketPrediction, PaperBet, PaperExclusion, PaperOrder, ShadowBet
from backend.markets import calibration, polymarket
from backend.markets.service import predict_market
from tests.conftest import login_client
from tests.test_orders import _second
from tests.test_portfolio import NOW, book, make_market, make_prediction  # noqa: F401  (fixture)


async def _shadows(s):
    return (await s.execute(select(ShadowBet).order_by(ShadowBet.created_at))).scalars().all()


# ---------- 1. Shadow bets ----------

async def test_a_bet_blocked_only_by_a_filter_is_followed_and_settled(db, book, monkeypatch):
    monkeypatch.setattr(settings, "LONGSHOT_MIN_PRICE", 0.40)     # the 36¢ ask becomes a "long shot"
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        p = await make_prediction(s, m)
        ev = await portfolio.apply_economics(s, m, p)
        assert ev.verdict == "NO" and [r["code"] for r in ev.reasons if r["blocking"]] == ["longshot"]
        assert (await s.execute(select(PaperBet))).first() is None
        [sh] = await _shadows(s)
        assert (sh.filter, sh.side, sh.status) == ("longshot", "YES", "open") and sh.shares > 0
        assert sh.avg_price == pytest.approx(0.36, abs=0.01)
        # The same forecast again: still one open shadow bet per market and filter
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert len(await _shadows(s)) == 1

        m.closed, m.resolved_yes, m.resolution, m.last_trading_price = True, True, "yes", 0.99
        await s.commit()
        assert await shadow.settle(s) == 1
        await s.refresh(sh)
        assert sh.status == "won" and sh.pnl == pytest.approx(sh.shares - sh.outlay)
        [row] = await shadow.report(s)
        assert (row["filter"], row["n"], row["won"], row["open"]) == ("longshot", 1, 1, 0)
        assert row["pnl"] > 0 and row["roi"] > 0 and row["avg_move"] > 0     # the filter cost money here


async def test_no_shadow_when_something_else_blocks_too(db, book, monkeypatch):
    monkeypatch.setattr(settings, "LONGSHOT_MIN_PRICE", 0.40)
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35, liquidity=100)          # illiquid: not a filter we measure
        p = await make_prediction(s, m)
        ev = await portfolio.apply_economics(s, m, p)
        assert {"longshot", "illiquid"} <= {r["code"] for r in ev.reasons if r["blocking"]}
        assert await _shadows(s) == []


async def test_shadow_for_second_opinion_and_objective_evidence(db, book, jev_client, monkeypatch):
    jev_client(noul_value=0.8, score_value=3.0)
    _second(monkeypatch, 0.30)
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        await predict_market(s, m)
        assert [x.filter for x in await _shadows(s)] == ["second_opinion"]

    monkeypatch.setattr(settings, "EVIDENCE_OBJECTIVE_HALF", 1.5)   # one article: the facts say weak evidence
    _second(monkeypatch, 0.6)
    async with SessionLocal() as s:
        m = await make_market(s, mid="m-2", yes=0.35, event="e2")
        p = await predict_market(s, m)
        assert p.signal == "HOLD" and p.jev_evidence_strength > p.evidence_strength
        sh = [x for x in await _shadows(s) if x.market_id == "m-2"]
        assert [x.filter for x in sh] == ["objective_evidence"] and sh[0].prediction_id == p.id


async def test_shadow_for_the_closing_line_guard(db, book):
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        row = await portfolio.get_settings(s)
        row.paused_at = NOW
        await s.commit()
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert [x.filter for x in await _shadows(s)] == ["clv_guard"]
        # A category the guard excluded counts as the guard; one the user excluded does not
        row.paused_at = None
        s.add(PaperExclusion(kind="category", value="Politics", label="x", source="guard"))
        s.add(PaperExclusion(kind="category", value="Sport", label="mine"))
        await s.commit()
        m2 = await make_market(s, mid="m-2", yes=0.35, event="e2", category="Politics")
        m3 = await make_market(s, mid="m-3", yes=0.35, event="e3", category="Sport")
        await portfolio.update_market_category(s, m2)
        await portfolio.update_market_category(s, m3)
        await portfolio.apply_economics(s, m2, await make_prediction(s, m2))
        await portfolio.apply_economics(s, m3, await make_prediction(s, m3))
        assert sorted((x.market_id, x.filter) for x in await _shadows(s)) == [("m-2", "clv_guard"), ("m-fed", "clv_guard")]


async def test_shadow_api_summary_export_and_reset(db, book, monkeypatch):
    monkeypatch.setattr(settings, "LONGSHOT_MIN_PRICE", 0.40)
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
    async with login_client("viewer") as api:
        data = (await api.get("/portfolio/shadow")).json()
        assert data["summary"][0]["filter"] == "longshot" and data["bets"][0]["filter_label"]
        assert (await api.get("/portfolio")).json()["shadow"][0]["n"] == 1
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO((await api.get("/portfolio/export", headers={"X-Lang": "en"})).content))
        ws = wb["Shadow"]
        assert ws.max_row == 2 and [c.value for c in ws[1]][:3] == ["shadow_id", "filter", "filter_label"]
        assert any(str(r[0].value).startswith("shadow_longshot") for r in wb["Summary"].iter_rows())
    async with login_client("admin") as api:
        assert (await api.post("/portfolio/reset", json={"bankroll": 1000})).status_code == 200
        assert (await api.get("/portfolio/shadow")).json()["bets"] == []


# ---------- 2. Second-opinion scorecard ----------

async def test_second_opinion_scorecard_on_resolved_markets(db):
    async with SessionLocal() as s:
        # (price, Jev, second opinion, signal, outcome)
        cases = [(0.40, 0.70, 0.30, "BUY_YES", False),   # disagreed, the second opinion was right
                 (0.40, 0.70, 0.35, "BUY_YES", True),    # disagreed, Jev was right
                 (0.60, 0.30, 0.20, "BUY_NO", False),    # agreed
                 (0.50, 0.55, 0.50, "HOLD", True)]
        for i, (price, jev, second, signal, outcome) in enumerate(cases):
            s.add(Market(id=f"r{i}", question=f"R{i}?", closed=True, resolved_yes=outcome, resolution="yes" if outcome else "no"))
            s.add(MarketPrediction(market_id=f"r{i}", market_probability=price, model_probability=jev, evidence_strength=0.8,
                                   blended_probability=price, edge=0.1, signal=signal, second_opinion=second,
                                   second_opinion_provider="gemini", article_count=1))
        await s.commit()
        so = await calibration.second_opinion_summary(s)
    assert so["n"] == 4 and so["providers"] == {"gemini": 4}
    assert so["disagreements"] == 2 and so["jev_right"] == 1
    # Blocked bets: YES at 40¢ lost (−0.40), YES at 40¢ won (+0.60) → +0.10 per share
    assert so["blocked_result_per_share"] == pytest.approx(0.10)
    assert so["brier_second"] < so["brier_model"] and so["gain_second"]["markets"] == 4
    async with login_client("viewer") as api:
        assert (await api.get("/predictions/calibration")).json()["second_opinion"]["n"] == 4


# ---------- 3. Maker fills from the price history ----------

async def test_order_fills_when_the_history_traded_below_the_limit(db, book, monkeypatch):
    monkeypatch.setattr(settings, "PAPER_ORDER_MODE", "maker")
    monkeypatch.setattr(settings, "MAKER_FILL_FROM_HISTORY", True)
    calls = []
    points: list = []

    async def history(token, start, end, client=None, fidelities=None):
        calls.append((token, fidelities))
        return points
    monkeypatch.setattr(polymarket, "fetch_price_history", history)
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        m.best_bid, m.best_ask = 0.34, 0.36
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        order = (await s.execute(select(PaperOrder))).scalar_one()
        assert order.limit_price == pytest.approx(0.35)
        t0 = order.created_at = orders._now() - timedelta(hours=1)   # placed an hour ago
        order.expires_at = t0 + timedelta(hours=6)
        await s.commit()
        # Traded at the limit only: others may be ahead in the queue, not filled
        points[:] = [(t0 + timedelta(minutes=5), 0.35), (t0 + timedelta(minutes=9), 0.36)]
        assert (await orders.process_orders(s))["filled"] == 0
        assert calls[-1] == ("yes-fed", (1, 60))
        # A dip below the limit between two syncs: filled then, at the limit
        points[:] = [(t0 + timedelta(minutes=5), 0.36), (t0 + timedelta(minutes=12), 0.34), (t0 + timedelta(minutes=20), 0.37)]
        assert (await orders.process_orders(s))["filled"] == 1
        bet = (await s.execute(select(PaperBet))).scalar_one()
        assert bet.avg_price == pytest.approx(0.35) and bet.entry == "maker"
        assert bet.created_at == t0 + timedelta(minutes=12)


async def test_history_unavailable_falls_back_to_the_snapshot(db, book, monkeypatch):
    monkeypatch.setattr(settings, "PAPER_ORDER_MODE", "maker")
    monkeypatch.setattr(settings, "MAKER_FILL_FROM_HISTORY", True)

    async def broken(*a, **k):
        raise RuntimeError("CLOB unreachable")
    monkeypatch.setattr(polymarket, "fetch_price_history", broken)
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.35)
        m.best_bid, m.best_ask = 0.34, 0.36
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert (await orders.process_orders(s)) == {"filled": 0, "expired": 0, "cancelled": 0}
        m.best_ask = 0.35
        await s.commit()
        assert (await orders.process_orders(s))["filled"] == 1
