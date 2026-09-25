from datetime import datetime, timedelta, timezone

import pytest
from sqlalchemy import select

from backend.betting import portfolio
from backend.db.database import SessionLocal
from backend.db.models import (
    Article, Market, MarketArticleLink, MarketPrediction, PaperBet, ProcessedArticle, Source,
)
from backend.markets import polymarket
from backend.markets.forecast import compute_signal
from tests.conftest import login_client

NOW = datetime.now(timezone.utc)


@pytest.fixture
def book(monkeypatch):
    """Order books by token id; a missing token raises like an unreachable CLOB."""
    books = {
        "yes-fed": polymarket.OrderBook(asks=[(0.36, 500), (0.37, 1000), (0.40, 3000)], bids=[(0.34, 800)]),
        "no-fed": polymarket.OrderBook(asks=[(0.66, 2000)], bids=[(0.64, 800)]),
    }

    async def fetch(token_id, client=None):
        if token_id not in books:
            raise RuntimeError("CLOB unreachable")
        return books[token_id]
    monkeypatch.setattr(polymarket, "fetch_order_book", fetch)
    return books


async def make_market(s, mid="m-fed", yes=0.35, liquidity=50_000, days=30, event="fed-december", tokens=True, category="Economy"):
    m = Market(id=mid, question=f"Question {mid}?", event_slug=event, yes_price=yes, liquidity=liquidity,
               volume=1e6, end_date=NOW + timedelta(days=days),
               yes_token_id="yes-fed" if tokens else None, no_token_id="no-fed" if tokens else None)
    s.add(m)
    src = (await s.execute(select(Source))).scalars().first()
    if src is None:
        src = Source(name="Reuters", url="https://r.test/rss")
        s.add(src)
        await s.flush()
    a = Article(source_id=src.id, title=f"News for {mid}", url=f"https://n.test/{mid}", url_hash=f"h-{mid}", fetched_at=NOW)
    s.add(a)
    await s.flush()
    s.add(ProcessedArticle(article_id=a.id, category=category))
    s.add(MarketArticleLink(market_id=mid, article_id=a.id, similarity=0.8))
    await s.commit()
    return m


async def make_prediction(s, market, p_jev=0.8, evidence=0.75):
    sig = compute_signal(p_jev, market.yes_price, evidence)
    p = MarketPrediction(market_id=market.id, market_probability=market.yes_price, model_probability=p_jev,
                         evidence_strength=evidence, blended_probability=sig.blended_probability,
                         model_weight=sig.model_weight, edge=sig.edge, signal=sig.signal,
                         kelly_fraction=sig.kelly_fraction, article_count=1)
    s.add(p)
    await s.commit()
    return p


async def test_auto_bet_with_order_book_and_no_duplicates(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        p = await make_prediction(s, m)
        ev = await portfolio.apply_economics(s, m, p)
        assert ev.verdict == "GO" and ev.quote_source == "book"
        assert p.economics["verdict"] == "GO" and p.economics["stake"] == pytest.approx(40)  # 4% of 1000
        assert m.category == "Economy"
        bets = (await s.execute(select(PaperBet))).scalars().all()
        assert len(bets) == 1 and bets[0].side == "YES" and bets[0].avg_price == pytest.approx(0.36)
        # A second forecast on the same market does not stack a second bet (and the market cap is full)
        p2 = await make_prediction(s, m)
        ev2 = await portfolio.apply_economics(s, m, p2)
        assert ev2.verdict == "NO" and any(r["code"] == "exposure_cap" for r in p2.economics["reasons"])
        assert len((await s.execute(select(PaperBet))).scalars().all()) == 1
        led = await portfolio.ledger(s)
        assert led["cash"] == pytest.approx(960)


async def test_estimate_when_book_unavailable(db, book):
    async with SessionLocal() as s:
        m = await make_market(s, tokens=False)
        ev = await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert ev.quote_source == "estimate" and ev.notes


@pytest.mark.parametrize("kind", ["market", "event", "category"])
async def test_exclusions_block_auto_bets(db, book, kind):
    async with SessionLocal() as s:
        m = await make_market(s)
        value = {"market": "m-fed", "event": "fed-december", "category": "Economy"}[kind]
        await portfolio.add_exclusion(s, kind, value)
        ev = await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert ev.verdict == "GO"  # the evaluation is still computed and shown
        assert (await s.execute(select(PaperBet))).first() is None


async def test_auto_off_and_manual_bet(db, book):
    async with SessionLocal() as s:
        await portfolio.update_settings(s, auto_paper=False)
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        assert (await s.execute(select(PaperBet))).first() is None
    async with login_client("admin") as api:
        r = await api.post("/markets/m-fed/paper-bet")
        assert r.status_code == 200 and r.json()["placed_by"] == "manual"
        assert (await api.post("/markets/m-fed/paper-bet")).status_code == 409


async def test_settlement_pnl_curve_and_exclusion(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        m.yes_price = 0.5
        await s.commit()
        summ = await portfolio.summary(s)
        bet = (await s.execute(select(PaperBet))).scalar_one()
        assert summ["unrealized_pnl_mid"] == pytest.approx(bet.shares * 0.5 - 40)
        # Selling now: at the bid (no book synced: mid − half the default spread), minus the sale fee
        from backend.betting.fees import fee_per_share
        from backend.config import settings
        bid = 0.5 - settings.DEFAULT_SPREAD / 2
        assert summ["unrealized_pnl"] == pytest.approx(bet.shares * (bid - fee_per_share(bid, portfolio.fee_bps_of(m))) - 40)
        m.resolved_yes, m.closed = True, True
        await s.commit()
        assert await portfolio.settle_bets(s) == 1
        await s.refresh(bet)
        assert bet.status == "won" and bet.pnl == pytest.approx(bet.shares - 40)
        summ = await portfolio.summary(s)
        assert summ["counts"]["won"] == 1 and summ["hit_rate"] == 1.0
        assert summ["equity"] == pytest.approx(1000 + bet.pnl)
        assert [p["equity"] for p in summ["equity_curve"]] == [1000, pytest.approx(1000 + bet.pnl, abs=0.01)]

    async with login_client("admin") as api:
        assert (await api.post(f"/portfolio/bets/{bet.id}/exclude")).status_code == 200
        assert (await api.get("/portfolio")).json()["equity"] == pytest.approx(1000)
        r = await api.post(f"/portfolio/bets/{bet.id}/include")
        assert r.json()["status"] == "won"  # re-included and settled again
        assert (await api.get("/portfolio")).json()["equity"] == pytest.approx(1000 + bet.pnl)


async def test_losing_no_bet(db, book):
    book["no-fed"] = polymarket.OrderBook(asks=[(0.36, 2000)], bids=[(0.34, 500)])  # NO costs ~36¢ when YES is 65¢
    async with SessionLocal() as s:
        m = await make_market(s, yes=0.65)
        p = await make_prediction(s, m, p_jev=0.15)
        ev = await portfolio.apply_economics(s, m, p)
        assert ev.side == "NO" and ev.verdict in ("GO", "SMALL")
        m.resolved_yes = True
        await s.commit()
        await portfolio.settle_bets(s)
        bet = (await s.execute(select(PaperBet))).scalar_one()
        assert bet.status == "lost" and bet.pnl == pytest.approx(-(bet.stake + bet.fee))


async def test_portfolio_api_permissions_settings_reset(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
    async with login_client("viewer") as viewer:
        body = (await viewer.get("/portfolio")).json()
        assert body["settings"]["preset"] == "bilanciato" and len(body["presets"]) == 3
        assert body["counts"]["open"] == 1 and body["invested"] == pytest.approx(40)
        assert (await viewer.get("/portfolio/bets", params={"status": "open"})).json()[0]["side"] == "YES"
        eco = (await viewer.get("/markets/m-fed/economics", params={"preset": "prudente"})).json()
        assert eco["preset"]["key"] == "prudente" and eco["open_bet"] is not None
        assert (await viewer.put("/portfolio/settings", json={"preset": "aggressivo"})).status_code == 403
        assert (await viewer.post("/portfolio/exclusions", json={"kind": "market", "value": "m-fed"})).status_code == 403
    async with login_client("admin") as api:
        assert (await api.put("/portfolio/settings", json={"preset": "aggressivo", "auto_paper": False})).json()["preset"] == "aggressivo"
        r = await api.post("/portfolio/exclusions", json={"kind": "category", "value": "Crypto", "label": "Crypto"})
        assert r.status_code == 201
        ex = (await api.get("/portfolio/exclusions")).json()
        assert ex[0]["value"] == "Crypto"
        assert (await api.delete(f"/portfolio/exclusions/{ex[0]['id']}")).status_code == 204
        assert (await api.post("/portfolio/reset", json={"bankroll": 5})).status_code == 422
        assert (await api.post("/portfolio/reset", json={"bankroll": 2500, "preset": "prudente"})).status_code == 200
        body = (await api.get("/portfolio")).json()
        assert body["equity"] == 2500 and body["counts"]["open"] == 0 and body["profile"]["key"] == "prudente"


async def test_split_resolution_pays_half(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        bet = (await s.execute(select(PaperBet))).scalar_one()
        m.resolution, m.closed = "split", True
        await s.commit()
        assert await portfolio.settle_bets(s) == 1
        await s.refresh(bet)
        assert bet.status == "void" and bet.payout == pytest.approx(bet.shares * 0.5)
        assert bet.pnl == pytest.approx(bet.shares * 0.5 - bet.stake - bet.fee)
        summ = await portfolio.summary(s)
        assert summ["counts"]["void"] == 1 and summ["hit_rate"] is None
        assert summ["equity"] == pytest.approx(1000 + bet.pnl)


async def test_fee_rate_from_category_unless_the_market_is_fee_free(db, book, monkeypatch):
    from backend.betting import fees
    async with SessionLocal() as s:
        m = await make_market(s)
        m.category = "Crypto"
        assert (await portfolio.build_quote(m, "YES")).fee_bps == 0      # conftest: CLOB says fee-free

        async def enabled(token_id, client=None):
            return True
        monkeypatch.setattr(fees, "fees_enabled", enabled)
        assert (await portfolio.build_quote(m, "YES")).fee_bps == pytest.approx(700)
        m.category = "Foreign Affairs"
        assert (await portfolio.build_quote(m, "YES")).fee_bps == 0

        async def unknown(token_id, client=None):
            return None
        monkeypatch.setattr(fees, "fees_enabled", unknown)
        m.category = "Politics"
        assert (await portfolio.build_quote(m, "YES")).fee_bps == pytest.approx(400)


def test_parse_fee_switch():
    from backend.betting.fees import parse_fee_enabled
    assert parse_fee_enabled({"base_fee": 0}) is False
    assert parse_fee_enabled({"base_fee": 1000}) is True
    assert parse_fee_enabled({"other": 1}) is None and parse_fee_enabled("x") is None


async def test_closing_line_value(db, book):
    from backend.markets import service as market_service
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        bet = (await s.execute(select(PaperBet))).scalar_one()
        # Still trading at 0.45: provisional move, no closing line yet
        m.yes_price = 0.45
        await s.commit()
        summ = await portfolio.summary(s)
        assert summ["clv"]["n"] == 0
        assert summ["clv_open"]["avg"] == pytest.approx(0.45 - bet.avg_price, abs=1e-4)
        # The sync records the last trading price, and keeps it when the market closes at 1
        data = polymarket.PolymarketMarket(id=m.id, question=m.question, slug=None, event_slug=m.event_slug,
                                           description=None, end_date=m.end_date, yes_price=0.52, volume=1e6,
                                           liquidity=5e4, active=True, closed=False, resolved_yes=None)
        market_service._apply_market(m, data)
        data.closed, data.yes_price, data.resolved_yes, data.resolution = True, 1.0, True, "yes"
        market_service._apply_market(m, data)
        await s.commit()
        assert m.last_trading_price == 0.52 and m.yes_price == 1.0
        summ = await portfolio.summary(s)
        assert summ["clv"] == {"n": 1, "avg": pytest.approx(0.52 - bet.avg_price, abs=1e-4), "share_positive": 1.0}
    async with login_client("viewer") as api:
        bets = (await api.get("/portfolio/bets")).json()
        assert bets[0]["clv"] == pytest.approx(0.52 - bet.avg_price, abs=1e-4)
        cal = (await api.get("/predictions/calibration")).json()
        # The BUY_YES signal was made at 0.35 and the market closed at 0.52
        assert cal["signal_clv"]["n"] == 1 and cal["signal_clv"]["avg"] == pytest.approx(0.17)
        assert cal["resolved_markets"] == 1 and cal["gain_blended"]["markets"] == 1


async def test_economics_returns_the_plan(db, book):
    async with SessionLocal() as s:
        await portfolio.update_settings(s, auto_paper=False)
        m = await make_market(s)
        await make_prediction(s, m)
    async with login_client("viewer") as api:
        data = (await api.get("/markets/m-fed/economics")).json()
        plan = data["strategy"]
        assert plan["action"] == "BUY" and plan["side"] == "YES"
        assert plan["orders"][0]["limit"] == data["evaluation"]["limit_price"]
        assert plan["levels"]["sell_above"] > plan["orders"][0]["limit"]
        assert plan["pros"] and plan["cons"] and plan["exit"]
    async with login_client("admin") as api:
        assert (await api.post("/markets/m-fed/paper-bet")).status_code == 200
        plan = (await api.get("/markets/m-fed/economics")).json()["strategy"]
        assert plan["action"] == "HOLD" and plan["orders"][0]["type"] == "sell"


async def test_open_bet_is_sold_when_the_price_reaches_the_estimate(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        bet = (await s.execute(select(PaperBet))).scalar_one()
        # Still below the target: the plan is to hold, with a sale target
    async with login_client("viewer") as api:
        item = (await api.get("/portfolio/bets")).json()[0]
        assert item["plan"]["action"] == "HOLD" and item["plan"]["sell_above"] > 0.5
        target = item["plan"]["sell_above"]
    # The price climbs to the target: the review sells into the bids
    book["yes-fed"] = polymarket.OrderBook(asks=[(target + 0.01, 5000)], bids=[(target, 5000)])
    async with SessionLocal() as s:
        m = await s.get(Market, "m-fed")
        m.yes_price, m.best_bid, m.best_ask = target, target, target + 0.01
        await s.commit()
        assert await portfolio.review_open_bets(s) == 1
        bet = await s.get(PaperBet, bet.id)
        assert bet.status == "sold" and bet.exit_price == pytest.approx(target)
        # Proceeds net of the sale fee, minus what was paid (stake and both fees)
        assert bet.pnl == pytest.approx(bet.shares * target - bet.stake - bet.fee)
        assert bet.pnl > 0 and "raggiunto la stima" in bet.exit_reason
        summ = await portfolio.summary(s)
        assert summ["counts"]["sold"] == 1 and summ["hit_rate"] == 1.0
        assert summ["realized_pnl"] == pytest.approx(bet.pnl)


async def test_auto_sell_off_and_manual_sell(db, book):
    async with SessionLocal() as s:
        await portfolio.update_settings(s, auto_sell=False)
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        m.yes_price = 0.9
        book["yes-fed"] = polymarket.OrderBook(asks=[(0.91, 5000)], bids=[(0.9, 5000)])
        await s.commit()
        assert await portfolio.review_open_bets(s) == 0   # automatic selling is off
        bet = (await s.execute(select(PaperBet))).scalar_one()
    async with login_client("viewer") as api:
        assert (await api.post(f"/portfolio/bets/{bet.id}/sell")).status_code == 403
    async with login_client("admin") as api:
        r = await api.post(f"/portfolio/bets/{bet.id}/sell")
        assert r.status_code == 200 and r.json()["status"] == "sold" and r.json()["exit_reason"] == "Venduta a mano"
        assert (await api.post(f"/portfolio/bets/{bet.id}/sell")).status_code == 409
        assert (await api.put("/portfolio/settings", json={"auto_sell": True})).json()["auto_sell"] is True


async def test_excluding_and_readmitting_a_closed_bet_keeps_its_history(db, book):
    """A lost bet excluded and readmitted keeps its closing time and result (equity curve)."""
    closed_at = NOW - timedelta(days=2)
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        m.resolved_yes, m.closed = False, True
        await s.commit()
        await portfolio.settle_bets(s)
        bet = (await s.execute(select(PaperBet))).scalar_one()
        bet.settled_at = closed_at
        await s.commit()
        pnl = bet.pnl
        assert bet.status == "lost" and pnl < 0
    async with login_client("admin") as api:
        assert (await api.post(f"/portfolio/bets/{bet.id}/exclude")).json()["status"] == "excluded"
        summ = (await api.get("/portfolio")).json()
        assert summ["equity"] == pytest.approx(1000) and summ["realized_pnl"] == pytest.approx(0)
        assert summ["counts"]["lost"] == 0 and len(summ["equity_curve"]) == 1
        assert (await api.post(f"/portfolio/bets/{bet.id}/include")).json()["status"] == "lost"
        summ = (await api.get("/portfolio")).json()
    async with SessionLocal() as s:
        bet = await s.get(PaperBet, bet.id)
        assert bet.settled_at == closed_at and bet.pnl == pytest.approx(pnl)
    assert summ["equity"] == pytest.approx(1000 + pnl)
    assert datetime.fromisoformat(summ["equity_curve"][-1]["t"]) == closed_at


async def test_excluding_and_readmitting_a_sold_bet_keeps_it_sold(db, book):
    """A bet sold before resolution is readmitted as sold, not as an open position."""
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        bet = (await s.execute(select(PaperBet))).scalar_one()
        assert await portfolio.sell_bet(s, bet, m, "Venduta a mano")
        await s.refresh(bet)
        sold = (bet.exit_price, bet.fee, bet.pnl, bet.settled_at)
    async with login_client("admin") as api:
        before = (await api.get("/portfolio")).json()
        await api.post(f"/portfolio/bets/{bet.id}/exclude")
        assert (await api.post(f"/portfolio/bets/{bet.id}/include")).json()["status"] == "sold"
        after = (await api.get("/portfolio")).json()
        assert (await api.get("/portfolio/bets", params={"status": "open"})).json() == []
    async with SessionLocal() as s:
        bet = await s.get(PaperBet, bet.id)
        assert (bet.exit_price, bet.fee, bet.pnl, bet.settled_at) == sold
    assert after["equity"] == pytest.approx(before["equity"]) and after["invested"] == pytest.approx(0)
    assert after["counts"]["sold"] == 1 and after["counts"]["open"] == 0


async def test_outcome_without_event_forecast_uses_its_own_forecast(db, book):
    """A position on an outcome of a multi-outcome event with no event forecast (placed before
    multi-outcome support) falls back to the market's own forecast: it gets an exit plan."""
    async with SessionLocal() as s:
        m = await make_market(s, mid="m-outcome")
        m.multi_event_id = "ev-1"
        await s.commit()
        p = await make_prediction(s, m)
        s.add(PaperBet(market_id=m.id, prediction_id=p.id, side="YES", shares=100, avg_price=0.36, stake=36.0, fee=0.5,
                       p_side=0.5, p_conservative=0.45, expected_profit=5.0, preset="bilanciato", status="open", placed_by="auto"))
        await s.commit()
        assert (await portfolio.latest_prediction(s, m)).id == p.id
    async with login_client("admin") as api:
        bets = (await api.get("/portfolio/bets", params={"status": "open"})).json()
    assert bets[0]["plan"]["action"] in ("HOLD", "SELL") and bets[0]["plan"]["sell_above"] is not None


async def test_open_bet_is_valued_at_the_bid_it_would_sell_at(db, book):
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        m.yes_price, m.best_bid, m.best_ask = 0.5, 0.47, 0.53
        await s.commit()
        bet = (await s.execute(select(PaperBet))).scalar_one()
        from backend.betting.fees import fee_per_share
        assert portfolio.mark_value(bet, m) == pytest.approx(bet.shares * (0.47 - fee_per_share(0.47, portfolio.fee_bps_of(m))))
        assert portfolio.mid_value(bet, m) == pytest.approx(bet.shares * 0.5)
    async with login_client("viewer") as api:
        row = (await api.get("/portfolio/bets", params={"status": "open"})).json()[0]
    assert row["bid_price"] == 0.47 and row["current_price"] == 0.5
    assert row["unrealized_pnl"] < row["unrealized_pnl_mid"]


async def test_opportunities_flag_exposure_reasons_that_no_longer_hold(db, book):
    """An opportunity blocked because the market's exposure was full: after the position is sold
    the stored reason no longer holds, and the API says so."""
    async with SessionLocal() as s:
        m = await make_market(s)
        p = await make_prediction(s, m)
        await portfolio.apply_economics(s, m, p)            # opens the position
        p2 = await make_prediction(s, m)
        await portfolio.apply_economics(s, m, p2)           # market cap now full: exposure_cap
        await s.refresh(p2)
        assert any(r["code"] == "exposure_cap" and r["blocking"] for r in p2.economics["reasons"])
        bet = (await s.execute(select(PaperBet))).scalar_one()
    async with login_client("viewer") as api:
        opp = (await api.get("/predictions/opportunities")).json()[0]
        assert "stale" not in (opp["prediction"]["economics"] or {})
    async with SessionLocal() as s:
        m = await s.get(Market, "m-fed")
        bet = await s.get(PaperBet, bet.id)
        assert await portfolio.sell_bet(s, bet, m, "Venduta a mano")
    async with login_client("viewer") as api:
        opp = (await api.get("/predictions/opportunities")).json()[0]
    assert opp["prediction"]["economics"]["stale"] == ["exposure_cap"]


async def test_readmitting_a_sold_bet_excluded_by_the_old_code(db, book):
    """Bets excluded before the fix lost pnl and closing time: readmitted sold, the result is
    rebuilt from the sale (fee already includes the sale fee) and the summary still works."""
    async with SessionLocal() as s:
        m = await make_market(s)
        await portfolio.apply_economics(s, m, await make_prediction(s, m))
        bet = (await s.execute(select(PaperBet))).scalar_one()
        assert await portfolio.sell_bet(s, bet, m, "Venduta a mano")
        await s.refresh(bet)
        pnl = bet.pnl
        bet.status, bet.payout, bet.pnl, bet.settled_at = "excluded", None, None, None   # the old exclusion
        await s.commit()
    async with login_client("admin") as api:
        assert (await api.post(f"/portfolio/bets/{bet.id}/include")).json()["status"] == "sold"
        summ = (await api.get("/portfolio")).json()
    assert summ["realized_pnl"] == pytest.approx(pnl) and len(summ["equity_curve"]) == 2


async def test_export_workbook_and_csv(db, book):
    """The export has every bet (open, sold, excluded) with its entry forecast and current value,
    the summary, the equity curve and the exclusions; readable by a viewer; in the chosen language."""
    import csv
    import io
    from openpyxl import load_workbook
    async with SessionLocal() as s:
        m = await make_market(s)
        p = await make_prediction(s, m)
        await portfolio.apply_economics(s, m, p)
        sold = (await s.execute(select(PaperBet))).scalar_one()
        assert await portfolio.sell_bet(s, sold, m, "Venduta a mano")
        m2 = await make_market(s, mid="m-btc", event="btc")
        await portfolio.apply_economics(s, m2, await make_prediction(s, m2))
        await portfolio.add_exclusion(s, "category", "Crypto", "Crypto")
        await s.refresh(sold)
    async with login_client("viewer") as api:
        r = await api.get("/portfolio/export", params={"format": "xlsx", "lang": "en"})
        summ = (await api.get("/portfolio")).json()
        r_csv = await api.get("/portfolio/export", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'attachment; filename="news-markets-portfolio-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Bets", "Orders", "Equity", "Exclusions", "Columns"]

    summary = {row[0]: row[2] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["bankroll"] == 1000 and summary["count_sold"] == 1 and summary["count_open"] == 1
    assert summary["realized_pnl"] == pytest.approx(summ["realized_pnl"])
    assert summary["total_value"] == pytest.approx(summ["total_value"])

    rows = list(wb["Bets"].iter_rows(values_only=True))
    bets = [dict(zip(rows[0], r)) for r in rows[1:]]
    assert len(bets) == 2
    s_row = next(b for b in bets if b["status"] == "sold")
    o_row = next(b for b in bets if b["status"] == "open")
    assert s_row["bet_id"] == str(sold.id) and s_row["pnl"] == pytest.approx(sold.pnl)
    assert s_row["exit_price"] == pytest.approx(sold.exit_price) and s_row["exit_reason"] == "Venduta a mano"
    assert s_row["jev_probability"] == pytest.approx(0.8) and s_row["signal"] == "BUY_YES" and s_row["verdict"] in ("GO", "SMALL")
    # UTC without time zone (Excel has none), to the millisecond
    assert abs(s_row["settled_at"] - sold.settled_at.astimezone(timezone.utc).replace(tzinfo=None)) < timedelta(milliseconds=1)
    assert o_row["market_id"] == "m-btc" and o_row["current_value"] > 0 and o_row["plan_action"] in ("HOLD", "SELL")
    assert o_row["unrealized_pnl"] == pytest.approx(o_row["current_value"] - o_row["outlay"])
    assert len(list(wb["Equity"].iter_rows(min_row=2))) == 2          # start + the sale
    assert [c.value for c in wb["Exclusions"][2]][:2] == ["category", "Crypto"]
    columns = {r[1]: r[2] for r in wb["Columns"].iter_rows(min_row=2, values_only=True) if r and r[1]}
    assert columns["clv"].startswith("Closing line value") and set(rows[0]) <= set(columns)

    assert r_csv.status_code == 200 and r_csv.headers["content-type"].startswith("text/csv")
    text = r_csv.content.decode("utf-8-sig")
    lines = list(csv.DictReader(io.StringIO(text)))
    assert len(lines) == 2 and {l["status"] for l in lines} == {"open", "sold"}
    assert float(next(l for l in lines if l["status"] == "sold")["pnl"]) == pytest.approx(sold.pnl)


async def test_export_sheet_names_follow_the_language(db):
    from openpyxl import load_workbook
    import io
    async with login_client("viewer") as api:
        r = await api.get("/portfolio/export", params={"format": "xlsx"}, headers={"X-Lang": "it"})
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Riepilogo", "Scommesse", "Ordini", "Capitale", "Esclusioni", "Colonne"]
